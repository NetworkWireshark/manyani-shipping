#!/usr/bin/env python3
"""MANYANI SHIPPING — local SQLite API + static UI."""
from __future__ import annotations

import csv
from order_excel import parse_order
import order_export
import order_tracking
from contextlib import closing
import hashlib
import hmac
import secrets
from http.cookies import SimpleCookie
import glob
import io
import json
import os
import re
import sqlite3
import threading
import tempfile
import time
import zipfile
from datetime import date, datetime, timedelta
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, quote, unquote, urlparse
import urllib.error
import urllib.request
ROOT = os.path.dirname(os.path.abspath(__file__))
PUBLIC = os.path.join(ROOT, "public")


def load_dotenv(path: str) -> None:
    if not os.path.isfile(path):
        return
    with open(path, encoding="utf-8") as handle:
        for line in handle:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            if key and key not in os.environ:
                os.environ[key] = value.strip()


load_dotenv(os.path.join(ROOT, ".env"))
DB_PATH = os.environ.get("DB_PATH", os.path.join(ROOT, "data", "marine_ops.sqlite"))
BACKUP_LOCAL = os.environ.get("BACKUP_LOCAL", os.path.join(ROOT, "data", "backups"))
BACKUP_SHARE = os.environ.get(
    "BACKUP_SHARE", os.path.join("E:\\Partage", "IT", "MANYANI SHIPPING", "backups")
)
PORT = int(os.environ.get("PORT", "80"))
HOST = os.environ.get("HOST", "0.0.0.0")
SITE = os.environ.get("SITE", "erica-bignoniaceous-tyson.ngrok-free.dev")
LAN_IP = os.environ.get("LAN_IP", "192.168.100.100")
TZ = datetime.now().astimezone().tzinfo

AUTH_PATH = os.path.join(os.path.dirname(DB_PATH), "auth.json")
AUTH_LOCK = threading.Lock()
SESSIONS = {}
LOGIN_ATTEMPTS = {}
SESSION_SECONDS = 8 * 3600


def auth_account():
    with AUTH_LOCK:
        if not os.path.exists(AUTH_PATH):
            os.makedirs(os.path.dirname(AUTH_PATH), exist_ok=True)
            password = os.environ.get("AUTH_PASSWORD") or secrets.token_urlsafe(15)
            salt = secrets.token_hex(16)
            account = {"username": os.environ.get("AUTH_USERNAME", "admin"), "salt": salt,
                       "hash": hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), 600000).hex()}
            with open(AUTH_PATH, "w", encoding="utf-8") as handle:
                json.dump(account, handle)
            with open(os.path.join(os.path.dirname(AUTH_PATH), "login-credentials.txt"), "w", encoding="utf-8") as handle:
                handle.write(f"Username: {account['username']}\nPassword: {password}\n")
        with open(AUTH_PATH, encoding="utf-8") as handle:
            stored = json.load(handle)
            return stored if "users" in stored else {"users": [dict(stored, role="editor")]}


STATUSES = ("LIVRÉ", "EN COURS DE LIVRAISON", "EN ATTENTE")
LIVREURS = (
    "AISSA TAYEB",
    "BOUFTAIM MUSTAPHA",
    "AIT AHMED MUSTAPHA",
    "MBAREEK GHAZLANE",
    "AZIZ",
    "AIT AHMAD AMINE",
    "JEDALI YASSINE",
    "ABDELGHAFOUR MANYANI",
    "MOHAMMED AMINE BOULEKHLEF",
    "SAID AMATERAZ",
    "TAOUIL MAHMOUDI",
)
LIVREUR_ALIASES = {
    "AISSA": "AISSA TAYEB",
    "AISA": "AISSA TAYEB",
    "AISSSA": "AISSA TAYEB",
    "AIISA": "AISSA TAYEB",
    "ISSA": "AISSA TAYEB",
    "TAYEB AISSA": "AISSA TAYEB",
    "AISSA TAYEB": "AISSA TAYEB",
    "BOUFTAIM": "BOUFTAIM MUSTAPHA",
    "BOUFAIM": "BOUFTAIM MUSTAPHA",
    "BOUFTIM": "BOUFTAIM MUSTAPHA",
    "BOFTAIM": "BOUFTAIM MUSTAPHA",
    "BOUFTAIN": "BOUFTAIM MUSTAPHA",
    "MOSTAFA BOUFTAIM": "BOUFTAIM MUSTAPHA",
    "MUSTAPHA BOUFTAIM": "BOUFTAIM MUSTAPHA",
    "MOSTAPHA": "BOUFTAIM MUSTAPHA",
    "MOSTAFA": "BOUFTAIM MUSTAPHA",
    "BOUFTAIM MUSTAPHA": "BOUFTAIM MUSTAPHA",
    "AITAHMED": "AIT AHMED MUSTAPHA",
    "AIT AHMED": "AIT AHMED MUSTAPHA",
    "AIT AHMED MUSTAPHA": "AIT AHMED MUSTAPHA",
    "MBAREK": "MBAREEK GHAZLANE",
    "MBAREEK": "MBAREEK GHAZLANE",
    "MBAREK RHAZLANE": "MBAREEK GHAZLANE",
    "MBAREK GHEZLANI": "MBAREEK GHAZLANE",
    "MBAREK GHAZLANI": "MBAREEK GHAZLANE",
    "MBAREEK GHAZLANE": "MBAREEK GHAZLANE",
    "AZIZ": "AZIZ",
    "AZIZ HOURI": "AZIZ",
    "JEDALI": "JEDALI YASSINE",
    "JADALI": "JEDALI YASSINE",
    "JADALI YASSINE": "JEDALI YASSINE",
    "YASSINE": "JEDALI YASSINE",
    "YASSINE JADALI": "JEDALI YASSINE",
    "YASSINE JADALLI": "JEDALI YASSINE",
    "JEDALI YASSINE": "JEDALI YASSINE",
    "ABDELGHAFOUR": "ABDELGHAFOUR MANYANI",
    "ABDELGHAFOYR": "ABDELGHAFOUR MANYANI",
    "ABDELGAHFOUR": "ABDELGHAFOUR MANYANI",
    "ABDELAGHOUR": "ABDELGHAFOUR MANYANI",
    "AVDELGHAFOUR": "ABDELGHAFOUR MANYANI",
    "ABDLGHFOUR": "ABDELGHAFOUR MANYANI",
    "ABDLGHAFOUR": "ABDELGHAFOUR MANYANI",
    "ABDO": "ABDELGHAFOUR MANYANI",
    "ABDELGHAFOUR MANYANI": "ABDELGHAFOUR MANYANI",
    "AIT AHMAD AMINE": "AIT AHMAD AMINE",
    "AITAHMAD AMINE": "AIT AHMAD AMINE",
    "MOHAMMED AMINE BOULEKHLEF": "MOHAMMED AMINE BOULEKHLEF",
    "BOULEKHLEF": "MOHAMMED AMINE BOULEKHLEF",
    "YOUNESSEE": "YOUNESSE",
    "YOUNESSR": "YOUNESSE",
    "SAED": "SAID AMATERAZ",
    "SAID": "SAID AMATERAZ",
    "SAID AMATERAZ": "SAID AMATERAZ",
    "MR TAOUIL": "TAOUIL MAHMOUDI",
    "TAOUIL": "TAOUIL MAHMOUDI",
    "TAOUIL MAHMOUDI": "TAOUIL MAHMOUDI",
    "EXTERIEUR": "EXTERNE",
    "BOUWISIR HASSAN": "HASSAN BOUWISIR",
    "HASSAN BOUWISIR": "HASSAN BOUWISIR",
    "MAHMOUDI MOSTAPHA": "MUSTAPHA MAHMOUDI",
    "MUSTAPHA MAHMOUDI": "MUSTAPHA MAHMOUDI",
    "MAHMOUDI(TAOUIL)": "TAOUIL MAHMOUDI",
    "MAHMOUDI (TAOUIL)": "TAOUIL MAHMOUDI",
}
SKIP_LIVREUR = {"LIVREUR", "PAR EMAIL", "EX", "—"}
PORTS = (
    "TANGER VILLE",
    "TANGER-MED",
    "JORF LASFAR",
    "CASABLANCA",
    "AGADIR",
    "SAFI",
    "MOHAMMEDIA",
    "NADOR",
    "KENITRA",
)
PORT_ALIASES = {
    "CASA": "CASABLANCA",
    "CASABLANCA": "CASABLANCA",
    "JORF": "JORF LASFAR",
    "JORF LASFAR": "JORF LASFAR",
    "TANGER": "TANGER-MED",
    "TANGIER": "TANGER-MED",
    "TANGRE": "TANGER-MED",
    "TANGER MED": "TANGER-MED",
    "TANGER-MED": "TANGER-MED",
    "TANGER VILLE": "TANGER VILLE",
    "SAFI ATLANTIQUE": "SAFI",
    "SAFI": "SAFI",
}


def norm_port(value: str | None) -> str:
    raw = re.sub(r"\s+", " ", (value or "").strip())
    raw = re.sub(r"\s*\([^)]*\)\s*", " ", raw).strip()
    if not raw:
        return ""
    return PORT_ALIASES.get(raw.upper(), raw)


def norm_imo(value: str | None) -> str:
    digits = re.sub(r"\D", "", str(value or ""))
    return digits if len(digits) == 7 else ""


def ship_key(name: str) -> str:
    key = re.sub(r"\s+", " ", (name or "").strip().upper())
    return re.sub(r"^(?:M/V|M\.V\.|MV|M/T|MT)\s+", "", key)


NAV_STATUS = {
    0: "En route",
    1: "Au mouillage",
    2: "Non maître de sa manœuvre",
    3: "Manœuvre restreinte",
    4: "Limité par le tirant d'eau",
    5: "À quai",
    6: "Échoué",
    7: "En pêche",
    8: "Sous voiles",
    15: "Non défini",
}
_AIS_CACHE: dict[str, tuple[float, object]] = {}
_AIS_LOCK = threading.Lock()


class AisError(Exception):
    def __init__(self, message: str, status: int = 502):
        super().__init__(message)
        self.status = status


def imo_valid(value: str | None) -> bool:
    imo = norm_imo(value)
    if len(imo) != 7:
        return False
    total = sum(int(imo[i]) * (7 - i) for i in range(6))
    return total % 10 == int(imo[6])


def _group_text(item: dict, field: str) -> str:
    for group_name in ("static", "registry"):
        text = str((item.get(group_name) or {}).get(field) or "").strip()
        if text:
            return text
    return ""


def _imo_of(item: dict) -> str:
    for group_name in ("static", "registry"):
        imo = norm_imo((item.get(group_name) or {}).get("imo"))
        if imo:
            return imo
    return ""


def _name_of(item: dict) -> str:
    name = _group_text(item, "name")
    if name:
        return name
    return str((item.get("match") or {}).get("matched_name") or "").strip()


def exact_vessels(results: list, name: str = "", imo: str = "") -> list:
    want_imo = norm_imo(imo)
    want_name = ship_key(name)
    found = []
    for item in results or []:
        match = item.get("match") or {}
        if match.get("kind") != "exact":
            continue
        if want_imo:
            if _imo_of(item) == want_imo:
                found.append(item)
            continue
        if match.get("field") != "name":
            continue
        names = {
            ship_key(match.get("matched_name") or ""),
            ship_key((item.get("registry") or {}).get("name") or ""),
            ship_key((item.get("static") or {}).get("name") or ""),
        }
        if want_name and want_name in names:
            found.append(item)
    return found


def prefer_live(items: list) -> dict | None:
    if not items:
        return None

    def rank(item: dict):
        pos = item.get("position") or {}
        seen = pos.get("ts") or (item.get("registry") or {}).get("last_seen") or ""
        return (1 if pos.get("live") else 0, seen)

    return max(items, key=rank)


def choose_vessel(results: list, name: str = "", imo: str = "") -> tuple[dict | None, list]:
    """One ship when the IMO or the exact name identifies it. Otherwise the exact list."""
    exact = exact_vessels(results, name, imo)
    if len(exact) <= 1:
        return (exact[0] if exact else None), []
    imos = {_imo_of(item) for item in exact if _imo_of(item)}
    if len(imos) == 1 and all(_imo_of(item) for item in exact):
        return prefer_live(exact), []
    return None, exact


SHIP_TYPES = (
    (30, 30, "Pêche"),
    (31, 32, "Remorquage"),
    (33, 34, "Travaux"),
    (35, 35, "Militaire"),
    (36, 36, "Voilier"),
    (37, 37, "Plaisance"),
    (40, 49, "Grande vitesse"),
    (50, 50, "Pilote"),
    (51, 51, "Sauvetage"),
    (52, 52, "Remorqueur"),
    (53, 55, "Service portuaire"),
    (58, 58, "Médical"),
    (60, 69, "Passagers"),
    (70, 79, "Cargo"),
    (80, 89, "Tanker"),
    (90, 99, "Autre"),
)
PLAUSIBILITY_LABEL = {
    "high": "Élevée",
    "medium": "Moyenne",
    "low": "Faible",
    "unknown": "Non évaluée",
}
PLAUSIBILITY_FLAGS = {
    "implausible_speed": "Vitesse impossible",
    "unreconciled_track": "Trajectoire incohérente",
    "invalid_mmsi": "MMSI invalide",
    "speed_mismatch": "Vitesse différente du déplacement",
    "course_mismatch": "Route différente du déplacement",
    "moving_while_moored": "Signale à l'arrêt mais se déplace",
}


def _first_number(*values):
    for value in values:
        if value is None or value == "":
            continue
        try:
            number = float(value)
        except (TypeError, ValueError):
            continue
        if number != number:
            continue
        return number
    return None


def _json_number(value):
    number = _first_number(value)
    if number is None:
        return None
    return int(number) if number.is_integer() else number


def _ship_type(*values):
    for value in values:
        code = _json_number(value)
        if code is None or int(code) <= 0:
            continue
        return int(code)
    return None


def ship_type_label(code) -> str:
    if code is None:
        return ""
    for start, end, label in SHIP_TYPES:
        if start <= int(code) <= end:
            return label
    return ""


def _metres(*values):
    for value in values:
        number = _first_number(value)
        if number is None or number <= 0:
            continue
        return int(number) if number.is_integer() else number
    return None


def vessel_public(item: dict, attribution: str = "") -> dict:
    static = item.get("static") or {}
    reg = item.get("registry") or {}
    pos = item.get("position") or {}
    try:
        nav = int(pos["nav_status"]) if pos.get("nav_status") is not None else None
    except (TypeError, ValueError):
        nav = None
    heading = pos.get("heading")
    if heading == 511:
        heading = None
    license_id = pos.get("license")
    if license_id is None:
        license_id = static.get("license", reg.get("license"))
    ship_type = _ship_type(static.get("type"), reg.get("type"))
    plausibility = str(pos.get("plausibility") or "")
    accuracy = pos.get("pos_accuracy")
    flags = []
    for flag in pos.get("flags") or []:
        label = PLAUSIBILITY_FLAGS.get(str(flag))
        if label and label not in flags:
            flags.append(label)
    return {
        "found": True,
        "name": _name_of(item),
        "imo": _imo_of(item),
        "mmsi": item.get("mmsi") or "",
        "callsign": (static.get("callsign") or reg.get("callsign") or "") or "",
        "dest": (static.get("dest") or "") or "",
        "eta": (static.get("eta") or "") or "",
        "flag": item.get("flag_country") or "",
        "type": ship_type,
        "type_label": ship_type_label(ship_type),
        "length": _metres(static.get("length"), reg.get("length")),
        "beam": _metres(static.get("beam"), reg.get("beam")),
        "draught": _metres(static.get("draught")),
        "lat": pos.get("lat"),
        "lon": pos.get("lon"),
        "sog": pos.get("sog"),
        "cog": pos.get("cog"),
        "heading": heading,
        "rot": _json_number(pos.get("rot")),
        "pos_accuracy": accuracy if isinstance(accuracy, bool) else None,
        "nav_status": nav,
        "nav_label": NAV_STATUS.get(nav, ""),
        "plausibility": plausibility if plausibility in PLAUSIBILITY_LABEL else "",
        "plausibility_label": PLAUSIBILITY_LABEL.get(plausibility, ""),
        "flags": flags,
        "live": bool(pos.get("live")),
        "ts": pos.get("ts") or "",
        "license": "" if license_id is None else str(license_id),
        "attribution": attribution or "Données AIS fournies par Pelyr",
    }


def pelyr_get(query_path: str, ttl: float = 120):
    key = os.environ.get("PELYR_KEY", "").strip()
    if not key:
        raise AisError("Clé AIS absente")
    now = time.time()
    with _AIS_LOCK:
        hit = _AIS_CACHE.get(query_path)
        if hit and hit[0] > now:
            return hit[1]
    req = urllib.request.Request(
        "https://api.pelyr.com" + query_path,
        headers={
            "Authorization": "Bearer " + key,
            "Accept": "application/json",
            "User-Agent": "ManyaniShipping/1.0",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=12) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        if exc.code == 429:
            raise AisError("Trop de recherches navire. Réessayez dans une minute.", 429)
        if exc.code in (401, 403):
            raise AisError("Accès AIS refusé", exc.code)
        if exc.code in (400, 422):
            raise AisError("Recherche AIS invalide", exc.code)
        raise AisError("Recherche AIS indisponible", exc.code)
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError):
        raise AisError("Recherche AIS indisponible")
    with _AIS_LOCK:
        _AIS_CACHE[query_path] = (now + ttl, payload)
        if len(_AIS_CACHE) > 300:
            for old in sorted(_AIS_CACHE, key=lambda item: _AIS_CACHE[item][0])[:80]:
                _AIS_CACHE.pop(old, None)
    return payload


def pelyr_attribution(license_id) -> str:
    if license_id is None or license_id == "":
        return "Données AIS fournies par Pelyr"
    try:
        data = pelyr_get("/v1/sources", ttl=3600)
    except AisError:
        return "Données AIS fournies par Pelyr"
    want = str(license_id)
    for source in data.get("sources") or []:
        if str(source.get("id")) == want and source.get("attribution"):
            return str(source["attribution"])
    return "Données AIS fournies par Pelyr"


def ais_lookup_vessel(name: str = "", imo: str = "") -> dict:
    name = re.sub(r"\s+", " ", (name or "").strip())
    imo = norm_imo(imo)
    if imo and not imo_valid(imo):
        imo = ""
    if not imo and len(ship_key(name)) < 3:
        return {"found": False, "name": name, "imo": "", "candidates": []}
    if imo:
        path = "/v1/vessels/search?imo=" + quote(imo)
    else:
        path = "/v1/vessels/search?name=" + quote(ship_key(name)) + "&max=10"
    try:
        data = pelyr_get(path, ttl=90)
    except AisError as exc:
        if exc.status in (400, 422):
            return {"found": False, "name": name, "imo": imo, "candidates": []}
        raise
    chosen, others = choose_vessel(data.get("results") or [], name, imo)
    if not chosen:
        candidates = []
        for item in others[:8]:
            candidates.append({
                "name": _name_of(item),
                "imo": _imo_of(item),
                "mmsi": item.get("mmsi") or "",
                "flag": item.get("flag_country") or "",
            })
        return {
            "found": False,
            "name": name,
            "imo": imo,
            "candidates": candidates,
            "attribution": "Données AIS fournies par Pelyr",
        }
    pos = chosen.get("position") or {}
    static = chosen.get("static") or {}
    reg = chosen.get("registry") or {}
    license_id = pos.get("license")
    if license_id is None:
        license_id = static.get("license", reg.get("license"))
    return vessel_public(chosen, pelyr_attribution(license_id))


def marinetraffic_ship_id(mmsi: str = "", imo: str = "") -> str:
    """MarineTraffic ship id, the value that selects the vessel on the live map."""
    mmsi_digits = re.sub(r"\D", "", str(mmsi or ""))
    imo_digits = norm_imo(imo)
    if len(mmsi_digits) == 9:
        kind, term = "MMSI", mmsi_digits
    elif imo_valid(imo_digits):
        kind, term = "IMO", imo_digits
    else:
        return ""
    cache_key = "mt-ship:" + kind + ":" + term
    now = time.time()
    with _AIS_LOCK:
        hit = _AIS_CACHE.get(cache_key)
        if hit and hit[0] > now:
            return str(hit[1] or "")
    ship_id = ""
    url = "https://www.marinetraffic.com/en/global_search/search?term=" + quote(term)
    req = urllib.request.Request(
        url,
        headers={
            "Accept": "application/json, text/javascript, */*; q=0.01",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36",
            "X-Requested-With": "XMLHttpRequest",
            "Referer": "https://www.marinetraffic.com/en/ais/home",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=8) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
        for item in payload.get("results") or []:
            if str(item.get("type") or "").upper() != kind:
                continue
            if str(item.get("value") or "") != term:
                continue
            found = str(item.get("id") or "")
            if found.isdigit():
                ship_id = found
                break
    except (urllib.error.URLError, TimeoutError, json.JSONDecodeError, OSError):
        ship_id = ""
    with _AIS_LOCK:
        _AIS_CACHE[cache_key] = (now + (7 * 24 * 3600 if ship_id else 120), ship_id)
    return ship_id


def remember_imo(rec_id: str, stored: str, imo: str, role: str) -> None:
    if not (imo and rec_id.isdigit() and not stored and role == "editor"):
        return
    conn = db()
    try:
        conn.execute(
            "UPDATE deliveries SET imo=? WHERE id=? AND (imo IS NULL OR TRIM(imo)='')",
            (imo, int(rec_id)),
        )
        conn.commit()
    finally:
        conn.close()


def track_urls(name: str, imo: str = "") -> dict:
    n = (name or "").strip()
    id_ = norm_imo(imo)
    q = quote(n)
    mst = (
        f"https://www.myshiptracking.com/?imo={id_}"
        if id_
        else f"https://www.myshiptracking.com/vessels?name={q}"
    )
    return {"imo": id_, "mst": mst}


def split_livreurs(value) -> list[str]:
    if isinstance(value, (list, tuple)):
        parts = [str(p) for p in value]
    else:
        parts = re.split(r"[/,;]+|\s+ET\s+|\s+&\s+", str(value or ""), flags=re.I)
    out: list[str] = []
    for part in parts:
        name = re.sub(r"\s+", " ", part.strip())
        key = name.upper()
        if not name or key in SKIP_LIVREUR or re.fullmatch(r"\d+", name):
            continue
        if re.match(r"^\d{4}-\d{2}-\d{2}", name):
            continue
        name = LIVREUR_ALIASES.get(key, name)
        if name not in out:
            out.append(name)
    order = {n: i for i, n in enumerate(LIVREURS)}
    out.sort(key=lambda n: (order.get(n, 1000), n))
    return out


def join_livreurs(names: list[str]) -> str:
    return " / ".join(names)


def normalize_stored_livreurs(conn: sqlite3.Connection) -> None:
    for row in conn.execute("SELECT id, livreur FROM deliveries"):
        new = join_livreurs(split_livreurs(row["livreur"]))
        old = row["livreur"] or ""
        if new != old:
            conn.execute("UPDATE deliveries SET livreur = ? WHERE id = ?", (new, row["id"]))
CSV_HEADERS = [
    "VESSEL",
    "PORT",
    "AGENT",
    "ETB",
    "ETD",
    "QUAI",
    "COMMANDE",
    "STATUS",
    "DATE OF DELIVERY",
    "LIVREUR",
    "INVOICING DATE",
]
STATUS_ALIASES = {
    "LIVRE": "LIVRÉ",
    "LIVRÉ": "LIVRÉ",
    "LIVREE": "LIVRÉ",
    "DONE": "LIVRÉ",
    "DELIVERED": "LIVRÉ",
    "EN COURS DE LIVRAISON": "EN COURS DE LIVRAISON",
    "EN COURS": "EN COURS DE LIVRAISON",
    "IN DELIVERY": "EN COURS DE LIVRAISON",
    "EN ATTENTE": "EN ATTENTE",
    "WAITING": "EN ATTENTE",
    "PENDING": "EN ATTENTE",
}


def today() -> date:
    return datetime.now(TZ).date()


def week_bounds(d: date) -> tuple[date, date]:
    start = d - timedelta(days=d.weekday())
    return start, start + timedelta(days=6)


def norm_status(value: str | None) -> str:
    raw = (value or "").strip().upper()
    raw = raw.replace("É", "E").replace("È", "E")
    compact = re.sub(r"\s+", " ", raw)
    if compact in STATUS_ALIASES:
        return STATUS_ALIASES[compact]
    if "LIVR" in compact and "COURS" not in compact:
        return "LIVRÉ"
    if "COURS" in compact:
        return "EN COURS DE LIVRAISON"
    if "ATTENTE" in compact or compact in ("WAIT", "WAITING"):
        return "EN ATTENTE"
    return "EN ATTENTE"


def parse_date(value: str | None, year_hint: int | None = None) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat() if 2020 <= value.year <= 2040 else None
    text = str(value).strip()
    if not text:
        return None
    try:
        num = float(text.replace(",", "."))
        if 20000 < num < 80000:
            d = (datetime(1899, 12, 30) + timedelta(days=num)).date()
            return d.isoformat() if 2020 <= d.year <= 2040 else None
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d/%m/%y", "%d-%m-%y"):
        try:
            d = datetime.strptime(text[:10], fmt).date()
            if 2020 <= d.year <= 2040:
                return d.isoformat()
        except ValueError:
            continue
    glued = re.match(r"^(\d{1,2})/(\d{2})(20\d{2})$", text)
    if glued:
        return parse_date(f"{glued.group(1)}/{glued.group(2)}/{glued.group(3)}")
    piece = re.search(
        r"(\d{1,2})\s*[/\-.]\s*(\d{1,2})(?:\s*[/\-.]\s*(20\d{2}))?",
        text,
    )
    if piece:
        day, month, year = piece.group(1), piece.group(2), piece.group(3)
        if not year:
            found = re.search(r"(20\d{2})", text)
            if found:
                year = found.group(1)
            elif year_hint and 2020 <= int(year_hint) <= 2040:
                year = str(int(year_hint))
        if year:
            try:
                d = datetime(int(year), int(month), int(day)).date()
                if 2020 <= d.year <= 2040:
                    return d.isoformat()
            except ValueError:
                pass
    return None


def fr_date(value: str | None) -> str:
    return re.sub(
        r"\b(\d{4})-(\d{2})-(\d{2})\b",
        lambda m: f"{m.group(3)}/{m.group(2)}/{m.group(1)}",
        (value or "").strip(),
    )


def coerce_no(value) -> int | None:
    if value is None or value == "":
        return None
    try:
        n = int(float(str(value).strip()))
        return n if n > 0 else None
    except (TypeError, ValueError):
        return None


def year_from(delivery: str | None, invoicing: str | None, explicit: int | None) -> int | None:
    if explicit:
        try:
            y = int(explicit)
            if 2020 <= y <= 2040:
                return y
        except (TypeError, ValueError):
            pass
    for value in (delivery, invoicing):
        if value and re.match(r"^\d{4}-", value):
            y = int(value[:4])
            if 2020 <= y <= 2040:
                return y
    return today().year


def db_sync() -> dict:
    conn = db()
    ver = int(conn.execute("PRAGMA data_version").fetchone()[0])
    n = int(conn.execute("SELECT COUNT(*) FROM deliveries").fetchone()[0])
    mx = int(conn.execute("SELECT COALESCE(MAX(id), 0) FROM deliveries").fetchone()[0])
    conn.close()
    mtime = 0
    try:
        mtime = int(os.path.getmtime(DB_PATH) * 1000)
        wal = DB_PATH + "-wal"
        if os.path.isfile(wal):
            mtime = max(mtime, int(os.path.getmtime(wal) * 1000))
    except OSError:
        pass
    with closing(db()) as conn:
        revision = conn.execute("SELECT COALESCE(MAX(id),0) FROM order_item_history").fetchone()[0]
        files = conn.execute("SELECT COUNT(*), COALESCE(MAX(id),0) FROM order_files").fetchone()
    return {"v": ver, "n": n, "id": mx, "t": mtime, "articles": revision, "files": list(files)}


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def renumber_boards(conn: sqlite3.Connection) -> None:
    groups = [
        "category = 'Planning'",
        "category = 'TECLIM'",
        "IFNULL(category,'ops') NOT IN ('Planning','TECLIM') AND status = 'EN COURS DE LIVRAISON'",
        "IFNULL(category,'ops') NOT IN ('Planning','TECLIM') AND status = 'EN ATTENTE'",
    ]
    years = [r[0] for r in conn.execute("SELECT DISTINCT year FROM deliveries WHERE year IS NOT NULL ORDER BY year")]
    for y in years:
        groups.append(
            "IFNULL(category,'ops') NOT IN ('Planning','TECLIM') AND status = 'LIVRÉ' "
            f"AND year = {int(y)}"
        )
    groups.append(
        "IFNULL(category,'ops') NOT IN ('Planning','TECLIM') AND status = 'LIVRÉ' AND year IS NULL"
    )
    for where in groups:
        ids = [
            r[0]
            for r in conn.execute(
                f"""SELECT id FROM deliveries WHERE {where}
                    ORDER BY CASE WHEN date_of_delivery IS NULL OR date_of_delivery = '' THEN 1 ELSE 0 END,
                             date_of_delivery, id"""
            )
        ]
        for i, did in enumerate(ids, 1):
            conn.execute("UPDATE deliveries SET no = ? WHERE id = ?", (i, did))


def next_no(conn: sqlite3.Connection, category: str, status: str, year) -> int:
    cat = (category or "ops").lower()
    if cat == "planning":
        where, args = "category = 'Planning'", []
    elif cat == "teclim":
        where, args = "category = 'TECLIM'", []
    elif status == "LIVRÉ":
        where, args = "IFNULL(category,'ops') NOT IN ('Planning','TECLIM') AND status = 'LIVRÉ' AND year = ?", [year or today().year]
    else:
        where, args = "IFNULL(category,'ops') NOT IN ('Planning','TECLIM') AND status = ?", [status]
    n = conn.execute(f"SELECT COALESCE(MAX(no), 0) FROM deliveries WHERE {where}", args).fetchone()[0]
    return int(n) + 1


def init_db() -> None:
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    order_tracking.migrate(DB_PATH)
    conn = db()
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS deliveries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            no INTEGER,
            vessel TEXT NOT NULL,
            imo TEXT,
            port TEXT,
            agent TEXT,
            etb TEXT,
            etd TEXT,
            quai TEXT,
            commande TEXT,
            status TEXT NOT NULL,
            date_of_delivery TEXT,
            livreur TEXT,
            invoicing_date TEXT,
            category TEXT,
            year INTEGER
        )
        """
    )
    conn.execute("""CREATE TABLE IF NOT EXISTS order_files (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        delivery_id INTEGER NOT NULL REFERENCES deliveries(id) ON DELETE CASCADE,
        filename TEXT NOT NULL,
        content BLOB NOT NULL,
        created_at TEXT NOT NULL
    )""")
    order_columns = {row[1] for row in conn.execute("PRAGMA table_info(order_files)")}
    if "parsed_json" not in order_columns:
        conn.execute("ALTER TABLE order_files ADD COLUMN parsed_json TEXT")
    if "chargeur" not in order_columns:
        conn.execute("ALTER TABLE order_files ADD COLUMN chargeur TEXT")
    conn.execute("""CREATE TABLE IF NOT EXISTS load_history_clear (
        id INTEGER PRIMARY KEY CHECK(id = 1),
        last_event_id INTEGER NOT NULL DEFAULT 0,
        cleared_by TEXT NOT NULL,
        cleared_at TEXT NOT NULL
    )""")
    order_tracking.prime(conn)
    conn.execute("UPDATE deliveries SET livreur = 'AISSA TAYEB' WHERE TRIM(livreur) = 'AISSA'")
    conn.execute("UPDATE deliveries SET livreur = 'BOUFTAIM MUSTAPHA' WHERE TRIM(livreur) = 'BOUFTAIM'")
    normalize_stored_livreurs(conn)
    conn.execute("UPDATE deliveries SET port = 'CASABLANCA' WHERE UPPER(TRIM(port)) IN ('CASA', 'CASABLANCA')")
    conn.execute("UPDATE deliveries SET port = 'JORF LASFAR' WHERE UPPER(TRIM(port)) IN ('JORF', 'JORF LASFAR')")
    conn.execute(
        "UPDATE deliveries SET port = 'TANGER-MED' WHERE UPPER(REPLACE(TRIM(port), ' ', '')) IN ('TANGER', 'TANGERMED')"
    )
    conn.execute("UPDATE deliveries SET port = 'TANGER VILLE' WHERE UPPER(TRIM(port)) = 'TANGER VILLE'")
    cols = {row[1] for row in conn.execute("PRAGMA table_info(deliveries)")}
    if "imo" not in cols:
        conn.execute("ALTER TABLE deliveries ADD COLUMN imo TEXT")
    conn.commit()
    missing = conn.execute("SELECT COUNT(*) FROM deliveries WHERE no IS NULL").fetchone()[0]
    if missing:
        renumber_boards(conn)
        conn.commit()
    for field in ("etb", "etd"):
        for row_id, val in conn.execute(f"SELECT id, {field} FROM deliveries WHERE {field} LIKE '%-__-__%'"):
            new = fr_date(val)
            if new != val:
                conn.execute(f"UPDATE deliveries SET {field} = ? WHERE id = ?", (new, row_id))
    conn.commit()
    for row in conn.execute(
        "SELECT id, date_of_delivery, invoicing_date, etb, etd, year FROM deliveries"
    ):
        raw = (row["date_of_delivery"] or "").strip()
        if not raw or re.match(r"^\d{4}-\d{2}-\d{2}$", raw):
            continue
        hint = row["year"]
        extra = f"{row['etb'] or ''} {row['etd'] or ''} {row['invoicing_date'] or ''}"
        found = re.search(r"(20\d{2})", extra)
        if found:
            hint = int(found.group(1))
        new = parse_date(raw, year_hint=hint)
        if new and new != raw:
            conn.execute("UPDATE deliveries SET date_of_delivery = ? WHERE id = ?", (new, row["id"]))
    conn.commit()
    conn.close()


SEED = [
    (1, "MAVI VATAN", "CASABLANCA", "GLOBE MARINE", "A QUAI", "3 JRS", "POSTE 62", "PROVISION", "LIVRÉ", "2026-08-12", "AISSA TAYEB", "2026-08-13", "ops", 2026),
    (2, "STAR PEACE", "JORF LASFAR", "SEATRADE", "A QUAI", "24 H", "OCP", "BONDED STORES", "LIVRÉ", "2026-07-02", "BOUFTAIM MUSTAPHA", "2026-07-03", "ops", 2026),
    (3, "IRON LADY", "TANGER-MED", "ARKAS", "EN RADE", "2 JRS", "TC3", "STORES", "LIVRÉ", "2025-11-20", "AISSA TAYEB", "2025-11-21", "ops", 2025),
    (4, "LIVATYA", "CASABLANCA", "CMA CGM", "A QUAI", "3 JRS", "SOMAPORT", "PROVISION", "LIVRÉ", "2026-06-18", "BOUFTAIM MUSTAPHA", "2026-06-19", "TECLIM", 2026),
    (5, "MV NORLAND", "AGADIR", "COMATAM", "A QUAI", "48 H", "POSTE 3", "PIECES", "LIVRÉ", "2025-12-04", "AISSA TAYEB", "2025-12-05", "ops", 2025),
    (6, "MSC OSCAR", "CASABLANCA", "CMA CGM", "A QUAI", "3 JRS", "MASS CEREAL", "PRODUITS", "LIVRÉ", "2026-04-11", "BOUFTAIM MUSTAPHA", "2026-04-12", "ops", 2026),
    (7, "CMA CGM CONGO", "TANGER-MED", "CMA CGM", "A QUAI", "24 H", "TC3", "ROPES", "LIVRÉ", "2026-03-22", "AISSA TAYEB", "2026-03-23", "TECLIM", 2026),
    (8, "WADI ALARISH", "SAFI", "GLOBE MARINE", "A QUAI", "3 JRS", "OCP", "SERVICE", "LIVRÉ", "2025-09-30", "BOUFTAIM MUSTAPHA", "2025-10-01", "ops", 2025),
    (9, "BULK POLAND", "JORF LASFAR", "SEATRADE", "A QUAI", "5 JRS", "OCP", "PROVISION", "LIVRÉ", "2026-01-15", "AISSA TAYEB", "2026-01-16", "ops", 2026),
    (10, "ONE HANOI", "MOHAMMEDIA", "ARKAS", "A QUAI", "24 H", "POSTE 1", "STORES", "LIVRÉ", "2026-02-08", "BOUFTAIM MUSTAPHA", "2026-02-09", "ops", 2026),
    (11, "MAERSK SEVILLE", "CASABLANCA", "CMA CGM", "EN MOUVEMENT CASA", "3 JRS", "POSTE 62", "PROVISION", "EN COURS DE LIVRAISON", "2026-09-14", "AISSA TAYEB", "", "Planning", 2026),
    (12, "YM WITNESS", "JORF LASFAR", "GLOBE MARINE", "A QUAI", "24 H", "OCP", "BONDED STORES", "EN COURS DE LIVRAISON", "2026-09-15", "BOUFTAIM MUSTAPHA", "", "Planning", 2026),
    (13, "COSCO SHIPPING", "TANGER-MED", "ARKAS", "EN RADE", "2 JRS", "TC3", "STORES / PIECES", "EN COURS DE LIVRAISON", "2026-09-16", "AISSA TAYEB", "", "Planning", 2026),
    (14, "FEDERAL POLAR", "AGADIR", "COMATAM", "A QUAI", "48 H", "", "PRODUITS", "EN COURS DE LIVRAISON", "2026-09-18", "BOUFTAIM MUSTAPHA", "", "ops", 2026),
    (15, "BBC MAGELLAN", "CASABLANCA", "SEATRADE", "A QUAI", "3 JRS", "SOMAPORT", "ROPES", "EN COURS DE LIVRAISON", "2026-09-14", "", "", "Planning", 2026),
    (16, "HAFNIA LOTTE", "CASABLANCA", "GLOBE MARINE", "EN RADE", "3 JRS", "POSTE 62", "PROVISION", "EN ATTENTE", "2026-09-21", "AISSA TAYEB", "", "Planning", 2026),
    (17, "CLIPPER SANTIAGO", "JORF LASFAR", "SEATRADE", "EN MOUVEMENT JORF", "5 JRS", "OCP", "PIECES", "EN ATTENTE", "2026-09-22", "BOUFTAIM MUSTAPHA", "", "Planning", 2026),
    (18, "SANTA CATALINA", "TANGER-MED", "ARKAS", "EN RADE", "24 H", "TC3", "BONDED STORES", "EN ATTENTE", "2026-09-25", "", "", "TECLIM", 2026),
    (19, "NORDIC OASIS", "SAFI", "COMATAM", "A QUAI", "3 JRS", "OCP", "SERVICE", "EN ATTENTE", "2026-09-28", "AISSA TAYEB", "", "ops", 2026),
    (20, "SPAR TAURUS", "KENITRA", "GLOBE MARINE", "EN RADE", "2 JRS", "", "STORES", "EN ATTENTE", "2026-10-02", "BOUFTAIM MUSTAPHA", "", "Planning", 2026),
    (21, "ATLANTIC STAR", "CASABLANCA", "CMA CGM", "EN RADE", "3 JRS", "MASS CEREAL", "PROVISION / STORES", "EN ATTENTE", "2026-10-05", "", "", "TECLIM", 2026),
    (22, "MARAN GAS SPARTA", "MOHAMMEDIA", "ARKAS", "EN MOUVEMENT", "24 H", "", "PRODUITS", "EN ATTENTE", "2026-10-08", "", "", "ops", 2026),
]


def row_to_dict(row: sqlite3.Row) -> dict:
    return {k: row[k] for k in row.keys()}


def filters_from(qs: dict) -> dict:
    def one(key: str) -> str:
        vals = qs.get(key, [""])
        return (vals[0] if vals else "").strip()

    return {
        "q": one("q"),
        "port": one("port"),
        "status": one("status"),
        "agent": one("agent"),
        "livreur": one("livreur"),
        "category": one("category"),
        "year": one("year"),
        "date_from": one("from") or one("date_from"),
        "date_to": one("to") or one("date_to"),
        "no_invoice": one("no_invoice"),
        "attention": one("attention"),
        "window": one("window"),
        "sort": one("sort") or "date_of_delivery",
        "dir": one("dir") or "asc",
    }


def where_sql(f: dict) -> tuple[str, list]:
    clauses = ["1=1"]
    args: list = []
    if f["q"]:
        clauses.append(
            "(vessel LIKE ? OR port LIKE ? OR agent LIKE ? OR commande LIKE ? OR livreur LIKE ? OR quai LIKE ? OR imo LIKE ?)"
        )
        like = f"%{f['q']}%"
        args.extend([like, like, like, like, like, like, like])
    if f["port"]:
        clauses.append("port LIKE ?")
        args.append(f"%{f['port']}%")
    if f["status"]:
        clauses.append("status = ?")
        args.append(norm_status(f["status"]))
    if f["agent"]:
        clauses.append("agent LIKE ?")
        args.append(f"%{f['agent']}%")
    if f["livreur"]:
        if f["livreur"] == "—":
            clauses.append("(livreur IS NULL OR TRIM(livreur) = '')")
        else:
            clauses.append("livreur LIKE ?")
            args.append(f"%{f['livreur']}%")
    if f["category"]:
        clauses.append("LOWER(category) = LOWER(?)")
        args.append(f["category"])
    if f["year"]:
        clauses.append("year = ?")
        args.append(int(f["year"]))
    if f["date_from"]:
        dfrom = parse_date(f["date_from"])
        if dfrom:
            clauses.append("date_of_delivery >= ?")
            args.append(dfrom)
    if f["date_to"]:
        dto = parse_date(f["date_to"])
        if dto:
            clauses.append("date_of_delivery <= ?")
            args.append(dto)
    if str(f.get("no_invoice") or "").lower() in ("1", "true", "yes"):
        clauses.append("(invoicing_date IS NULL OR TRIM(invoicing_date) = '')")
    if f.get("attention") == "overdue":
        clauses.append("status != 'LIVRÉ' AND date_of_delivery < ? AND TRIM(IFNULL(date_of_delivery, '')) != ''")
        args.append(today().isoformat())
    elif f.get("attention") == "uninvoiced":
        clauses.append("status = 'LIVRÉ' AND (invoicing_date IS NULL OR TRIM(invoicing_date) = '')")
    if str(f.get("window") or "").lower() == "jour":
        d0 = today()
        d1 = d0 + timedelta(days=1)
        clauses.append("(date_of_delivery IN (?, ?) OR status = 'EN COURS DE LIVRAISON')")
        args.extend([d0.isoformat(), d1.isoformat()])
    return " AND ".join(clauses), args


SORTABLE = {
    "date_of_delivery": "date_of_delivery",
    "invoicing_date": "invoicing_date",
    "etb": "etb",
    "etd": "etd",
    "vessel": "vessel COLLATE NOCASE",
    "port": "port COLLATE NOCASE",
    "agent": "agent COLLATE NOCASE",
    "livreur": "livreur COLLATE NOCASE",
    "commande": "commande COLLATE NOCASE",
    "status": "status",
    "no": "COALESCE(no, id)",
    "id": "id",
}


def order_sql(f: dict) -> str:
    col = SORTABLE.get(f["sort"], "date_of_delivery")
    direction = "DESC" if f["dir"].lower() == "desc" else "ASC"
    # Default number order: active rows first, LIVRÉ (green) last.
    if f.get("sort") in ("", "no"):
        rank = (
            "CASE status "
            "WHEN 'EN COURS DE LIVRAISON' THEN 0 "
            "WHEN 'EN ATTENTE' THEN 1 "
            "ELSE 2 END"
        )
        return f"{rank}, {col} {direction}, id ASC"
    return f"{col} {direction}, id ASC"


def body_json(handler) -> dict:
    length = int(handler.headers.get("Content-Length") or 0)
    raw = handler.rfile.read(length) if length else b"{}"
    if not raw:
        return {}
    return json.loads(raw.decode("utf-8"))


def payload_to_record(data: dict, existing: dict | None = None) -> dict:
    base = existing or {}
    vessel = (data.get("vessel") or base.get("vessel") or "").strip()
    if not vessel:
        raise ValueError("Le navire (VESSEL) est obligatoire.")
    status = norm_status(data.get("status") or base.get("status") or "EN ATTENTE")
    try:
        year_hint = int(data.get("year") or (base or {}).get("year") or 0) or None
    except (TypeError, ValueError):
        year_hint = None
    delivery = parse_date(
        data.get("date_of_delivery") if "date_of_delivery" in data else base.get("date_of_delivery"),
        year_hint=year_hint,
    )
    invoicing = parse_date(
        data.get("invoicing_date") if "invoicing_date" in data else base.get("invoicing_date"),
        year_hint=year_hint,
    )
    for label, key, parsed in (
        ("DATE OF DELIVERY", "date_of_delivery", delivery),
        ("INVOICING DATE", "invoicing_date", invoicing),
    ):
        raw = data.get(key) if key in data else base.get(key)
        if raw is not None and str(raw).strip() and not parsed:
            raise ValueError(f"{label} : date invalide. Utilisez JJ/MM/AAAA.")
    year = year_from(delivery, invoicing, data.get("year") if "year" in data else base.get("year"))
    def field(name: str) -> str:
        if name in data:
            return (data.get(name) or "").strip()
        return (base.get(name) or "").strip() if base else ""

    rec = {
        "no": coerce_no(data.get("no") if "no" in data else base.get("no")),
        "vessel": vessel,
        "imo": norm_imo(field("imo")),
        "port": norm_port(field("port")),
        "agent": field("agent"),
        "etb": field("etb"),
        "etd": field("etd"),
        "quai": field("quai"),
        "commande": field("commande"),
        "status": status,
        "date_of_delivery": delivery or "",
        "livreur": join_livreurs(split_livreurs(data["livreur"] if "livreur" in data else field("livreur"))),
        "invoicing_date": invoicing or "",
        "category": field("category") or "Planning",
        "year": year,
    }
    return rec


def list_deliveries(f: dict) -> list[dict]:
    sql, args = where_sql(f)
    conn = db()
    rows = conn.execute(
        f"SELECT * FROM deliveries WHERE {sql} ORDER BY {order_sql(f)}", args
    ).fetchall()
    conn.close()
    return [row_to_dict(r) for r in rows]


def stats_for(f: dict) -> dict:
    rows = list_deliveries(f)
    d0 = today()
    w0, w1 = week_bounds(d0)
    by_port: dict[str, int] = {}
    by_livreur: dict[str, int] = {}
    counts = {"LIVRÉ": 0, "EN COURS DE LIVRAISON": 0, "EN ATTENTE": 0}
    today_n = 0
    week_n = 0
    teclim_n = 0
    teclim_open = 0
    teclim_livre = 0
    overdue_n = 0
    uninvoiced_n = 0
    for row in rows:
        st = row["status"]
        if st in counts:
            counts[st] += 1
        if (row["category"] or "").strip().lower() == "teclim":
            teclim_n += 1
            if st == "LIVRÉ":
                teclim_livre += 1
            else:
                teclim_open += 1
        port = (row["port"] or "—").strip() or "—"
        by_port[port] = by_port.get(port, 0) + 1
        names = split_livreurs(row["livreur"])
        if not names:
            by_livreur["—"] = by_livreur.get("—", 0) + 1
        for liv in names:
            by_livreur[liv] = by_livreur.get(liv, 0) + 1
        dod = row["date_of_delivery"] or ""
        if st != "LIVRÉ" and dod and dod < d0.isoformat():
            overdue_n += 1
        if st == "LIVRÉ" and not (row["invoicing_date"] or "").strip():
            uninvoiced_n += 1
        if dod == d0.isoformat():
            today_n += 1
        if dod and w0.isoformat() <= dod <= w1.isoformat():
            week_n += 1
    ports = sorted(by_port.items(), key=lambda x: (-x[1], x[0]))
    livres = sorted(by_livreur.items(), key=lambda x: (-x[1], x[0]))
    conn = db()
    years = [r[0] for r in conn.execute("SELECT DISTINCT year FROM deliveries WHERE year IS NOT NULL ORDER BY year DESC")]
    seen_ports = set(PORTS)
    extra_ports = []
    for r in conn.execute(
        "SELECT DISTINCT port FROM deliveries WHERE port IS NOT NULL AND TRIM(port) != '' ORDER BY port"
    ):
        name = norm_port(r[0])
        if name and name not in seen_ports:
            seen_ports.add(name)
            extra_ports.append(name)
    ports_all = list(PORTS) + extra_ports
    agents = [r[0] for r in conn.execute("SELECT DISTINCT agent FROM deliveries WHERE agent IS NOT NULL AND TRIM(agent) != '' ORDER BY agent")]
    seen = set(LIVREURS)
    extra = []
    for r in conn.execute(
        "SELECT DISTINCT livreur FROM deliveries WHERE livreur IS NOT NULL AND TRIM(livreur) != ''"
    ):
        for name in split_livreurs(r[0]):
            if name not in seen:
                seen.add(name)
                extra.append(name)
    extra.sort()
    livreurs = list(LIVREURS) + extra
    def n(sql: str, args: list | tuple = ()) -> int:
        return int(conn.execute(sql, args).fetchone()[0])

    boards = {
        "all": n("SELECT COUNT(*) FROM deliveries"),
        "planning": n(
            "SELECT COUNT(*) FROM deliveries WHERE LOWER(IFNULL(category,'')) = 'planning' AND status != 'LIVRÉ'"
        ),
        "teclim": n(
            "SELECT COUNT(*) FROM deliveries WHERE LOWER(IFNULL(category,'')) = 'teclim' AND status != 'LIVRÉ'"
        ),
        "teclim_livre": n(
            "SELECT COUNT(*) FROM deliveries WHERE LOWER(IFNULL(category,'')) = 'teclim' AND status = 'LIVRÉ'"
        ),
        "teclim_all": n(
            "SELECT COUNT(*) FROM deliveries WHERE LOWER(IFNULL(category,'')) = 'teclim'"
        ),
        "en_cours": n("SELECT COUNT(*) FROM deliveries WHERE status = 'EN COURS DE LIVRAISON'"),
        "attente": n("SELECT COUNT(*) FROM deliveries WHERE status = 'EN ATTENTE'"),
        "livre": n("SELECT COUNT(*) FROM deliveries WHERE status = 'LIVRÉ'"),
        "retard": n("SELECT COUNT(*) FROM deliveries WHERE status != 'LIVRÉ' AND date_of_delivery < ? AND TRIM(IFNULL(date_of_delivery, '')) != ''", (d0.isoformat(),)),
        "a_facturer": n("SELECT COUNT(*) FROM deliveries WHERE status = 'LIVRÉ' AND TRIM(IFNULL(invoicing_date, '')) = ''"),
        "jour": n(
            "SELECT COUNT(*) FROM deliveries WHERE date_of_delivery IN (?, ?) OR status = 'EN COURS DE LIVRAISON'",
            (d0.isoformat(), (d0 + timedelta(days=1)).isoformat()),
        ),
    }
    conn.close()
    newest = sorted(rows, key=lambda r: int(r.get("id") or 0), reverse=True)[:8]
    recent = sorted(
        newest,
        key=lambda r: ((r.get("date_of_delivery") or "0000-00-00"), int(r.get("id") or 0)),
        reverse=True,
    )
    return {
        "total": len(rows),
        "teclim": teclim_n,
        "teclim_open": teclim_open,
        "teclim_livre": teclim_livre,
        "livre": counts["LIVRÉ"],
        "en_cours": counts["EN COURS DE LIVRAISON"],
        "en_attente": counts["EN ATTENTE"],
        "overdue": overdue_n,
        "uninvoiced": uninvoiced_n,
        "today": today_n,
        "this_week": week_n,
        "week_range": {"from": w0.isoformat(), "to": w1.isoformat()},
        "today_date": d0.isoformat(),
        "tomorrow_date": (d0 + timedelta(days=1)).isoformat(),
        "by_port": [{"name": k, "count": v} for k, v in ports],
        "by_livreur": [{"name": k, "count": v} for k, v in livres],
        "years": years,
        "ports": ports_all,
        "agents": agents,
        "livreurs": livreurs,
        "statuses": list(STATUSES),
        "boards": boards,
        "recent": recent,
        "sync": db_sync(),
    }


def csv_bytes(rows: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\n")
    writer.writerow(CSV_HEADERS)
    for row in rows:
        writer.writerow(
            [
                row.get("vessel") or "",
                row.get("port") or "",
                row.get("agent") or "",
                fr_date(row.get("etb")),
                fr_date(row.get("etd")),
                row.get("quai") or "",
                row.get("commande") or "",
                row.get("status") or "",
                fr_date(row.get("date_of_delivery")),
                row.get("livreur") or "",
                fr_date(row.get("invoicing_date")),
            ]
        )
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def xlsx_bytes(rows: list[dict]) -> bytes:
    def cell(ref: str, value: str) -> str:
        safe = (
            (value or "")
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{safe}</t></is></c>'

    def col_letter(i: int) -> str:
        return chr(ord("A") + i)

    sheet_rows = []
    header_cells = "".join(cell(f"{col_letter(i)}1", h) for i, h in enumerate(CSV_HEADERS))
    sheet_rows.append(f'<row r="1">{header_cells}</row>')
    keys = [
        "vessel", "port", "agent", "etb", "etd", "quai", "commande",
        "status", "date_of_delivery", "livreur", "invoicing_date",
    ]
    for r_i, row in enumerate(rows, start=2):
        cells = "".join(
            cell(
                f"{col_letter(c)}{r_i}",
                fr_date(row.get(k)) if k in ("etb", "etd", "date_of_delivery", "invoicing_date") else str(row.get(k) or ""),
            )
            for c, k in enumerate(keys)
        )
        sheet_rows.append(f'<row r="{r_i}">{cells}</row>')
    sheet = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f'<sheetData>{"".join(sheet_rows)}</sheetData></worksheet>'
    )
    workbook = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="Livraisons" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        "</Relationships>"
    )
    wb_rels = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )
    ctypes = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        "</Types>"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", ctypes)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("xl/workbook.xml", workbook)
        zf.writestr("xl/_rels/workbook.xml.rels", wb_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet)
    return buf.getvalue()


HEADER_MAP = {
    "VESSEL": "vessel",
    "NAVIRE": "vessel",
    "IMO": "imo",
    "PORT": "port",
    "AGENT": "agent",
    "ETB": "etb",
    "ETD": "etd",
    "QUAI": "quai",
    "COMMANDE": "commande",
    "STATUS": "status",
    "STATUT": "status",
    "DATE OF DELIVERY": "date_of_delivery",
    "DATE OF DELIVERY ": "date_of_delivery",
    "DATE LIVRAISON": "date_of_delivery",
    "LIVREUR": "livreur",
    "INVOICING DATE": "invoicing_date",
    "DATE FACTURE": "invoicing_date",
    "CATEGORY": "category",
    "BOARD": "category",
    "YEAR": "year",
    "NO": "no",
}


def sniff_csv(text: str) -> list[dict]:
    sample = text[:2048]
    dialect = csv.excel
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,|\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    out = []
    for raw in reader:
        item = {}
        for key, val in raw.items():
            if key is None:
                continue
            mapped = HEADER_MAP.get(key.strip().upper())
            if mapped:
                item[mapped] = (val or "").strip()
        if item.get("vessel"):
            out.append(item)
    return out


def parse_xlsx(data: bytes) -> list[dict]:
    from xml.etree import ElementTree as ET

    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = {n.lower(): n for n in zf.namelist()}
        strings: list[str] = []
        ss = names.get("xl/sharedstrings.xml")
        if ss:
            root = ET.fromstring(zf.read(ss))
            for si in root.findall("m:si", ns):
                strings.append("".join(t.text or "" for t in si.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t")))
        sheet_name = next((n for n in zf.namelist() if n.startswith("xl/worksheets/sheet")), None)
        if not sheet_name:
            return []
        root = ET.fromstring(zf.read(sheet_name))
        rows: list[list[str]] = []
        for row in root.findall("m:sheetData/m:row", ns):
            cells: dict[int, str] = {}
            for c in row.findall("m:c", ns):
                ref = c.get("r", "A1")
                col = 0
                for ch in ref:
                    if ch.isalpha():
                        col = col * 26 + (ord(ch.upper()) - 64)
                    else:
                        break
                val = ""
                if c.get("t") == "s":
                    v = c.find("m:v", ns)
                    if v is not None and v.text:
                        val = strings[int(v.text)]
                elif c.get("t") == "inlineStr":
                    val = "".join(t.text or "" for t in c.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t"))
                else:
                    v = c.find("m:v", ns)
                    val = v.text if v is not None and v.text else ""
                cells[col - 1] = val
            if cells:
                width = max(cells) + 1
                rows.append([cells.get(i, "") for i in range(width)])
    if not rows:
        return []
    headers = [HEADER_MAP.get(h.strip().upper()) for h in rows[0]]
    out = []
    for row in rows[1:]:
        item = {}
        for i, key in enumerate(headers):
            if key and i < len(row):
                item[key] = row[i].strip()
        if item.get("vessel"):
            out.append(item)
    return out


def read_upload(handler) -> bytes:
    ctype = handler.headers.get("Content-Type", "")
    length = int(handler.headers.get("Content-Length") or 0)
    data = handler.rfile.read(length)
    if "multipart/form-data" in ctype:
        m = re.search(r"boundary=(.+)", ctype)
        if not m:
            return data
        boundary = m.group(1).strip().encode()
        parts = data.split(b"--" + boundary)
        for part in parts:
            if b"filename=" in part:
                _, _, body = part.partition(b"\r\n\r\n")
                return body.rstrip(b"\r\n")
        return b""
    return data


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=PUBLIC, **kwargs)

    def log_message(self, fmt: str, *args) -> None:
        print(f"[{datetime.now(TZ).strftime('%H:%M:%S')}] {fmt % args}")

    def end_headers(self):
        path = urlparse(self.path).path.lower()
        if path.endswith((".html", ".js", ".css")) or path in ("/", "/index.html"):
            self.send_header("Cache-Control", "no-store, max-age=0")
            self.send_header("Pragma", "no-cache")
        super().end_headers()

    def send_json(self, payload, status=200, extra_headers=None):
        raw = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(raw)))
        for key, value in extra_headers or ():
            self.send_header(key, value)
        self.end_headers()
        self.wfile.write(raw)

    def redirect(self, location: str, extra_headers=None):
        self.send_response(302)
        self.send_header("Location", location)
        self.send_header("Content-Length", "0")
        self.send_header("Cache-Control", "no-store")
        for key, value in extra_headers or ():
            self.send_header(key, value)
        self.end_headers()

    def send_file_bytes(self, payload: bytes, filename: str, mime: str):
        self.send_response(200)
        self.send_header("Content-Type", mime)
        self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)

    def qs(self) -> dict:
        return parse_qs(urlparse(self.path).query)

    def session(self):
        cookie = SimpleCookie()
        try:
            cookie.load(self.headers.get("Cookie", ""))
            token = cookie["manyani_session"].value if "manyani_session" in cookie else ""
        except Exception:
            return None
        with AUTH_LOCK:
            now = time.time()
            for key in list(SESSIONS):
                if SESSIONS[key]["expires"] <= now:
                    del SESSIONS[key]
            session = SESSIONS.get(token)
            if session and session.get("auth_hash"):
                try:
                    with open(AUTH_PATH, encoding="utf-8") as handle:
                        stored = json.load(handle)
                    users = stored["users"] if "users" in stored else [stored]
                    account = next((user for user in users if user["username"] == session["username"]), None)
                    if account is None or not hmac.compare_digest(session["auth_hash"], account["hash"]):
                        SESSIONS.pop(token, None)
                        return None
                except (OSError, ValueError, KeyError):
                    SESSIONS.pop(token, None)
                    return None
            return session

    def require_login(self):
        if self.session():
            return True
        if urlparse(self.path).path.startswith("/api/"):
            self.send_json({"error": "Connexion requise"}, 401)
        else:
            self.redirect("/login")
        return False

    def require_editor(self):
        if not self.require_login():
            return False
        if self.session().get("role") != "editor":
            self.send_json({"error": "Compte en lecture seule : modifications refusees"}, 403)
            return False
        return True

    def trusted_origin(self):
        origin = self.headers.get("Origin")
        return not origin or urlparse(origin).netloc == self.headers.get("Host")

    def session_cookie(self, token, age):
        secure = "; Secure" if self.headers.get("X-Forwarded-Proto", "").lower() == "https" else ""
        return ("Set-Cookie", f"manyani_session={token}; Path=/; HttpOnly; SameSite=Strict; Max-Age={age}{secure}")

    def do_HEAD(self):
        if not self.require_login():
            return
        return super().do_HEAD()

    def do_GET(self):
        parsed = urlparse(self.path)
        path = unquote(parsed.path)
        if path in ("/login", "/login.html"):
            if self.session():
                return self.redirect("/")
            self.path = "/login.html"
            return super().do_GET()
        if path == "/api/health":
            return self.send_json({"ok": True})
        if path in ("/logo.svg", "/logo.png"):
            return super().do_GET()
        if not self.require_login():
            return
        if path == "/api/session":
            return self.send_json({"username": self.session()["username"], "role": self.session().get("role", "viewer")})
        if path == "/api/accounts":
            return self.send_json([{"username": user["username"]} for user in auth_account()["users"] if user.get("role", "viewer") != "viewer"])
        if path == "/api/load-history":
            with closing(db()) as conn:
                clear = conn.execute("SELECT last_event_id, cleared_at FROM load_history_clear WHERE id=1").fetchone()
                cutoff = clear["last_event_id"] if clear else 0
                actions = order_tracking.LOAD_ACTIONS
                marks = ",".join("?" * len(actions))
                total = conn.execute(
                    f"SELECT COUNT(*) FROM order_item_history WHERE id>? AND action IN ({marks})",
                    (cutoff, *actions)).fetchone()[0]
                rows = conn.execute(f"""SELECT h.id, h.delivery_id, h.file_id, h.item_index,
                    h.article, h.reference, h.username, h.action, h.occurred_at,
                    h.before_value, h.after_value, h.original_quantity, h.after_quantity,
                    d.vessel, d.port, f.filename
                    FROM order_item_history h
                    LEFT JOIN deliveries d ON d.id=h.delivery_id
                    LEFT JOIN order_files f ON f.id=h.file_id
                    WHERE h.id>? AND h.action IN ({marks})
                    ORDER BY h.id DESC LIMIT 5000""", (cutoff, *actions)).fetchall()
            return self.send_json({"events": [dict(row) for row in rows], "total": total,
                "limited": total > len(rows), "cleared_at": clear["cleared_at"] if clear else None,
                "can_clear": self.session()["username"] == "admin"})
        if path == "/api/articles/availability":
            with closing(db()) as conn, conn:
                return self.send_json(order_tracking.availability_board(conn))
        if path == "/api/articles/loaded":
            with closing(db()) as conn, conn:
                return self.send_json(order_tracking.loaded_articles(conn))
        export_match = re.fullmatch(r"/api/deliveries/(\d+)/order-export\.(xlsx|pdf)", path)
        if export_match:
            with closing(db()) as conn, conn:
                delivery = conn.execute("SELECT * FROM deliveries WHERE id=?", (int(export_match[1]),)).fetchone()
                if not delivery:
                    return self.send_json({"error": "Livraison introuvable"}, 404)
                filters = {
                    "period": (self.qs().get("period") or ["all"])[0],
                    "user": (self.qs().get("user") or [""])[0][:100],
                    "status": (self.qs().get("status") or ["all"])[0],
                    "q": (self.qs().get("q") or [""])[0][:200],
                }
                payload, mime, filename = order_export.build(conn, delivery, filters, export_match[2])
                return self.send_file_bytes(payload, filename, mime)
        history_match = re.fullmatch(r"/api/deliveries/(\d+)/order-files/(\d+)/items/(\d+)/history", path)
        if history_match:
            with closing(db()) as conn:
                if not conn.execute("SELECT id FROM order_files WHERE delivery_id=? AND id=?", (int(history_match[1]),int(history_match[2]))).fetchone():
                    return self.send_json({"error": "Commande introuvable"}, 404)
                rows = conn.execute("SELECT * FROM order_item_history WHERE delivery_id=? AND file_id=? AND item_index=? ORDER BY id DESC",
                    (int(history_match[1]),int(history_match[2]),int(history_match[3]))).fetchall()
                return self.send_json([dict(row) for row in rows])
        order_match = re.fullmatch(r"/api/deliveries/(\d+)/order-files(?:/(\d+))?", path)
        if order_match:
            with closing(db()) as conn, conn:
                if not conn.execute("SELECT id FROM deliveries WHERE id=?", (int(order_match[1]),)).fetchone():
                    return self.send_json({"error": "Livraison introuvable"}, 404)
                if order_match[2]:
                    row = conn.execute("SELECT filename, content FROM order_files WHERE delivery_id=? AND id=?",
                                       (int(order_match[1]), int(order_match[2]))).fetchone()
                    if not row:
                        return self.send_json({"error": "Fichier introuvable"}, 404)
                    return self.send_file_bytes(row["content"], quote(row["filename"]),
                        "application/pdf" if row["filename"].lower().endswith(".pdf") else
                        "application/vnd.ms-excel" if row["filename"].lower().endswith(".xls") else
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                rows = conn.execute("SELECT id, filename, created_at, parsed_json, content, chargeur FROM order_files WHERE delivery_id=? ORDER BY id DESC",
                                    (int(order_match[1]),)).fetchall()
                files = []
                for row in rows:
                    entry = {key: row[key] for key in ("id", "filename", "created_at", "chargeur")}
                    try:
                        entry["order"] = order_tracking.enrich(conn, row["id"], order_tracking.read_order(row))
                        entry["loads"] = order_tracking.load_events(conn, row["id"])
                    except Exception:
                        entry["error"] = "Impossible de lire les articles de ce fichier."
                    files.append(entry)
                return self.send_json(files)
        if path == "/api/sync":
            return self.send_json(db_sync())
        if path == "/api/stats":
            return self.send_json(stats_for(filters_from(self.qs())))
        if path in ("/api/ais", "/api/ais-lookup"):
            qs = self.qs()
            name = (qs.get("name") or [""])[0].strip()[:80]
            asked = norm_imo((qs.get("imo") or [""])[0])
            rec_id = (qs.get("id") or [""])[0].strip()
            stored = ""
            if rec_id.isdigit():
                conn = db()
                row = conn.execute("SELECT vessel, imo FROM deliveries WHERE id=?", (int(rec_id),)).fetchone()
                conn.close()
                if row:
                    name = name or (row["vessel"] or "")
                    stored = norm_imo(row["imo"])
            try:
                card = ais_lookup_vessel(name, stored or asked)
            except AisError as exc:
                if path == "/api/ais":
                    return self.send_json({"error": str(exc)}, exc.status if exc.status < 500 else 502)
                card = {"found": False, "imo": ""}
            imo = stored or norm_imo(card.get("imo"))
            remember_imo(rec_id, stored, imo, self.session().get("role", ""))
            if path == "/api/ais":
                if imo and not card.get("imo"):
                    card["imo"] = imo
                card["name"] = card.get("name") or name
                if card.get("found"):
                    card["ship_id"] = marinetraffic_ship_id(card.get("mmsi") or "", card.get("imo") or "")
                return self.send_json(card)
            out = track_urls(name, imo)
            out["name"] = name
            if card.get("found"):
                out["position"] = card
            return self.send_json(out)
        if path == "/api/deliveries":
            return self.send_json(list_deliveries(filters_from(self.qs())))
        m = re.fullmatch(r"/api/deliveries/(\d+)", path)
        if m:
            conn = db()
            row = conn.execute("SELECT * FROM deliveries WHERE id=?", (int(m.group(1)),)).fetchone()
            conn.close()
            if not row:
                return self.send_json({"error": "Introuvable"}, 404)
            return self.send_json(row_to_dict(row))
        if path in ("/api/export.csv", "/api/export"):
            rows = list_deliveries(filters_from(self.qs()))
            return self.send_file_bytes(csv_bytes(rows), "manyani-shipping.csv", "text/csv; charset=utf-8")
        if path == "/api/export.xlsx":
            rows = list_deliveries(filters_from(self.qs()))
            return self.send_file_bytes(
                xlsx_bytes(rows),
                "manyani-shipping.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        if path == "/":
            self.path = "/index.html"
        return super().do_GET()

    def do_POST(self):
        parsed = urlparse(self.path)
        path = parsed.path
        if not self.trusted_origin():
            return self.send_json({"error": "Origine non autorisee"}, 403)
        if path == "/api/login":
            ip = self.client_address[0]
            now = time.time()
            with AUTH_LOCK:
                attempts = [stamp for stamp in LOGIN_ATTEMPTS.get(ip, []) if stamp > now - 300]
                LOGIN_ATTEMPTS[ip] = attempts
                if len(attempts) >= 10:
                    return self.send_json({"error": "Trop de tentatives. Reessayez dans 5 minutes."}, 429)
                attempts.append(now)
            try:
                payload = body_json(self)
                username, password = payload.get("username", ""), payload.get("password", "")
                if not isinstance(username, str) or not isinstance(password, str) or len(password) > 1024:
                    raise ValueError("Invalid credentials")
            except (ValueError, AttributeError):
                return self.send_json({"error": "Identifiants invalides"}, 400)
            users = auth_account()["users"]
            account = next((user for user in users if user["username"] == username), users[0])
            digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(account["salt"]), 600000).hex()
            if not (hmac.compare_digest(username.encode(), account["username"].encode()) and hmac.compare_digest(digest, account["hash"])):
                return self.send_json({"error": "Nom d'utilisateur ou mot de passe incorrect"}, 401)
            token = secrets.token_urlsafe(32)
            with AUTH_LOCK:
                LOGIN_ATTEMPTS.pop(ip, None)
                SESSIONS[token] = {"username": account["username"], "role": account.get("role", "viewer"), "auth_hash": account["hash"], "expires": now + SESSION_SECONDS}
            return self.send_json({"ok": True}, extra_headers=[self.session_cookie(token, SESSION_SECONDS)])
        if path == "/api/logout":
            cookie = SimpleCookie()
            try:
                cookie.load(self.headers.get("Cookie", ""))
                with AUTH_LOCK:
                    if "manyani_session" in cookie:
                        SESSIONS.pop(cookie["manyani_session"].value, None)
            except Exception:
                pass
            return self.send_json({"ok": True}, extra_headers=[self.session_cookie("", 0)])
        if path == "/api/change-password":
            if not self.require_login():
                return
            try:
                payload = body_json(self)
                old_password = payload.get("current_password")
                new_password = payload.get("new_password")
                if not isinstance(old_password, str) or not isinstance(new_password, str):
                    raise ValueError()
                if len(old_password) > 1024 or not 8 <= len(new_password) <= 128:
                    raise ValueError()
            except (ValueError, AttributeError):
                return self.send_json({"error": "Le nouveau mot de passe doit contenir entre 8 et 128 caracteres."}, 400)
            username = self.session()["username"]
            auth_account()  # Create the legacy account file if needed, before taking the lock.
            with AUTH_LOCK:
                with open(AUTH_PATH, encoding="utf-8") as handle:
                    stored = json.load(handle)
                users = stored["users"] if "users" in stored else [stored]
                account = next((user for user in users if user["username"] == username), None)
                if account is None:
                    return self.send_json({"error": "Compte introuvable"}, 404)
                digest = hashlib.pbkdf2_hmac("sha256", old_password.encode(), bytes.fromhex(account["salt"]), 600000).hex()
                if not hmac.compare_digest(digest, account["hash"]):
                    return self.send_json({"error": "Mot de passe actuel incorrect"}, 403)
                if hmac.compare_digest(old_password, new_password):
                    return self.send_json({"error": "Choisissez un nouveau mot de passe different"}, 400)
                salt = secrets.token_hex(16)
                account["salt"] = salt
                account["hash"] = hashlib.pbkdf2_hmac("sha256", new_password.encode(), bytes.fromhex(salt), 600000).hex()
                temporary = None
                try:
                    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=os.path.dirname(AUTH_PATH), prefix=".auth-", suffix=".tmp", delete=False) as handle:
                        temporary = handle.name
                        json.dump(stored, handle)
                        handle.flush()
                        os.fsync(handle.fileno())
                    os.replace(temporary, AUTH_PATH)
                finally:
                    if temporary and os.path.exists(temporary):
                        os.unlink(temporary)
                for token, session in list(SESSIONS.items()):
                    if session["username"] == username:
                        del SESSIONS[token]
            return self.send_json({"ok": True}, extra_headers=[self.session_cookie("", 0)])
        if path == "/api/load-history/clear":
            if not self.require_login():
                return
            if self.session()["username"] != "admin":
                return self.send_json({"error": "Action reservee au compte admin"}, 403)
            with closing(db()) as conn, conn:
                actions = order_tracking.LOAD_ACTIONS
                marks = ",".join("?" * len(actions))
                last_id = conn.execute(
                    f"SELECT COALESCE(MAX(id),0) FROM order_item_history WHERE action IN ({marks})", actions).fetchone()[0]
                conn.execute("""INSERT INTO load_history_clear(id,last_event_id,cleared_by,cleared_at)
                    VALUES(1,?,?,?) ON CONFLICT(id) DO UPDATE SET
                    last_event_id=excluded.last_event_id, cleared_by=excluded.cleared_by,
                    cleared_at=excluded.cleared_at""",
                    (last_id, "admin", datetime.now().astimezone().isoformat(timespec="seconds")))
            return self.send_json({"ok": True})
        if not self.require_editor():
            return
        order_match = re.fullmatch(r"/api/deliveries/(\d+)/order-files", path)
        if order_match:
            filename = (self.qs().get("filename") or [""])[0]
            filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
            try:
                length = int(self.headers.get("Content-Length", "0"))
            except ValueError:
                return self.send_json({"error": "Taille invalide"}, 400)
            if not filename.lower().endswith((".xls", ".xlsx", ".pdf")) or len(filename) > 200 or any(ord(c) < 32 for c in filename):
                return self.send_json({"error": "Choisissez un fichier Excel .xls, .xlsx ou PDF .pdf"}, 400)
            if not 0 < length <= 20 * 1024 * 1024:
                return self.send_json({"error": "Le fichier doit faire entre 1 octet et 20 Mo"}, 400)
            content = self.rfile.read(length)
            is_pdf = filename.lower().endswith(".pdf")
            valid_content = content.startswith(b"%PDF-") if is_pdf else (content.startswith(bytes.fromhex("d0cf11e0a1b11ae1")) or zipfile.is_zipfile(io.BytesIO(content)))
            if len(content) != length or not valid_content:
                return self.send_json({"error": "Fichier PDF invalide" if is_pdf else "Fichier Excel invalide"}, 400)
            try:
                order = parse_order(content)
            except Exception as exc:
                return self.send_json({"error": "Import impossible : " + str(exc)}, 400)
            with closing(db()) as conn, conn:
                delivery_id = int(order_match[1])
                if not conn.execute("SELECT id FROM deliveries WHERE id=?", (delivery_id,)).fetchone():
                    return self.send_json({"error": "Livraison introuvable"}, 404)
                payload = json.dumps(order, ensure_ascii=False)
                existing = conn.execute(
                    "SELECT id FROM order_files WHERE delivery_id=? AND lower(filename)=lower(?)",
                    (delivery_id, filename)).fetchone()
                if not existing:
                    existing = conn.execute(
                        "SELECT id FROM order_files WHERE delivery_id=? AND content=?",
                        (delivery_id, content)).fetchone()
                if existing:
                    # Same order again: keep the file, the Excel quantities and every load.
                    file_id = existing["id"]
                    conn.execute("UPDATE order_files SET filename=?, content=?, parsed_json=? WHERE id=?",
                        (filename, content, payload, file_id))
                    order_tracking.seed_items(conn, file_id, order)
                    conn.commit()
                    return self.send_json({"id": file_id, "filename": filename, "items": len(order["items"]), "replaced": True})
                cursor = conn.execute("INSERT INTO order_files (delivery_id, filename, content, created_at, parsed_json) VALUES (?, ?, ?, ?, ?)",
                    (delivery_id, filename, content, datetime.now(TZ).isoformat(), payload))
                order_tracking.seed_items(conn, cursor.lastrowid, order)
                conn.commit()
                return self.send_json({"id": cursor.lastrowid, "filename": filename, "items": len(order["items"])}, 201)
        if path == "/api/deliveries":
            try:
                rec = payload_to_record(body_json(self))
            except ValueError as exc:
                return self.send_json({"error": str(exc)}, 400)
            conn = db()
            if not rec.get("no"):
                rec["no"] = next_no(conn, rec["category"], rec["status"], rec["year"])
            cur = conn.execute(
                """
                INSERT INTO deliveries (
                    no, vessel, imo, port, agent, etb, etd, quai, commande, status,
                    date_of_delivery, livreur, invoicing_date, category, year
                ) VALUES (:no,:vessel,:imo,:port,:agent,:etb,:etd,:quai,:commande,:status,
                          :date_of_delivery,:livreur,:invoicing_date,:category,:year)
                """,
                rec,
            )
            conn.commit()
            row = conn.execute("SELECT * FROM deliveries WHERE id=?", (cur.lastrowid,)).fetchone()
            conn.close()
            return self.send_json(row_to_dict(row), 201)
        if path in ("/api/import.csv", "/api/import"):
            raw = read_upload(self)
            items: list[dict] = []
            if raw[:2] == b"PK":
                items = parse_xlsx(raw)
            else:
                text = raw.decode("utf-8-sig", errors="replace")
                items = sniff_csv(text)
            conn = db()
            inserted = 0
            errors = []
            for item in items:
                try:
                    rec = payload_to_record(item)
                    conn.execute(
                        """
                        INSERT INTO deliveries (
                            no, vessel, imo, port, agent, etb, etd, quai, commande, status,
                            date_of_delivery, livreur, invoicing_date, category, year
                        ) VALUES (:no,:vessel,:imo,:port,:agent,:etb,:etd,:quai,:commande,:status,
                                  :date_of_delivery,:livreur,:invoicing_date,:category,:year)
                        """,
                        rec,
                    )
                    inserted += 1
                except ValueError as exc:
                    errors.append(str(exc))
            conn.commit()
            conn.close()
            return self.send_json({"inserted": inserted, "errors": errors[:10]})
        return self.send_json({"error": "Not found"}, 404)

    def do_PUT(self):
        if not self.trusted_origin():
            return self.send_json({"error": "Origine non autorisee"}, 403)
        if not self.require_editor():
            return
        m = re.fullmatch(r"/api/deliveries/(\d+)", urlparse(self.path).path)
        if not m:
            return self.send_json({"error": "Not found"}, 404)
        conn = db()
        row = conn.execute("SELECT * FROM deliveries WHERE id=?", (int(m.group(1)),)).fetchone()
        if not row:
            conn.close()
            return self.send_json({"error": "Introuvable"}, 404)
        try:
            rec = payload_to_record(body_json(self), row_to_dict(row))
        except ValueError as exc:
            conn.close()
            return self.send_json({"error": str(exc)}, 400)
        rec["id"] = int(m.group(1))
        conn.execute(
            """
            UPDATE deliveries SET
                no=:no, vessel=:vessel, imo=:imo, port=:port, agent=:agent, etb=:etb, etd=:etd,
                quai=:quai, commande=:commande, status=:status,
                date_of_delivery=:date_of_delivery, livreur=:livreur,
                invoicing_date=:invoicing_date, category=:category, year=:year
            WHERE id=:id
            """,
            rec,
        )
        conn.commit()
        out = conn.execute("SELECT * FROM deliveries WHERE id=:id", rec).fetchone()
        conn.close()
        return self.send_json(row_to_dict(out))

    def do_PATCH(self):
        if not self.trusted_origin():
            return self.send_json({"error": "Origine non autorisee"}, 403)
        if not self.require_editor():
            return
        chargeur_match = re.fullmatch(r"/api/deliveries/(\d+)/order-files/(\d+)/chargeur", urlparse(self.path).path)
        if chargeur_match:
            try:
                chosen = body_json(self).get("chargeur", "")
                if not isinstance(chosen, str) or len(chosen) > 100:
                    raise ValueError("Chargeur invalide")
                chosen = chosen.strip()
                names = [user["username"] for user in auth_account()["users"] if user.get("role", "viewer") != "viewer"]
                if chosen and chosen not in names:
                    raise ValueError("Chargeur invalide")
            except (ValueError, TypeError) as exc:
                return self.send_json({"error": str(exc)}, 400)
            with closing(db()) as conn, conn:
                row = conn.execute("SELECT id FROM order_files WHERE delivery_id=? AND id=?", (int(chargeur_match[1]), int(chargeur_match[2]))).fetchone()
                if not row:
                    return self.send_json({"error": "Commande introuvable"}, 404)
                conn.execute("UPDATE order_files SET chargeur=? WHERE id=?", (chosen, int(chargeur_match[2])))
                return self.send_json({"chargeur": chosen})
        if urlparse(self.path).path == "/api/articles/availability":
            try:
                payload = body_json(self)
                with closing(db()) as conn, conn:
                    result = order_tracking.set_availability(
                        conn, int(payload.get("file_id")), int(payload.get("index")),
                        payload.get("availability"), payload.get("expected_version"),
                        self.session()["username"])
                    conn.commit()
                    return self.send_json(result)
            except order_tracking.ItemError as exc:
                return self.send_json({"error": str(exc)}, exc.status)
            except (ValueError, TypeError) as exc:
                return self.send_json({"error": str(exc)}, 400)
        item_match = re.fullmatch(r"/api/deliveries/(\d+)/order-files/(\d+)/items/(\d+)", urlparse(self.path).path)
        if item_match:
            try:
                payload = body_json(self)
                with closing(db()) as conn, conn:
                    result = order_tracking.change(conn,int(item_match[1]),int(item_match[2]),int(item_match[3]),payload,self.session()["username"])
                    conn.commit()
                    return self.send_json(result)
            except order_tracking.ItemError as exc:
                return self.send_json({"error": str(exc)},exc.status)
            except (ValueError,TypeError) as exc:
                return self.send_json({"error": str(exc)},400)
        m = re.fullmatch(r"/api/deliveries/(\d+)/status", urlparse(self.path).path)
        if not m:
            return self.send_json({"error": "Not found"}, 404)
        status = norm_status(body_json(self).get("status"))
        conn = db()
        row = conn.execute("SELECT * FROM deliveries WHERE id=?", (int(m.group(1)),)).fetchone()
        if not row:
            conn.close()
            return self.send_json({"error": "Introuvable"}, 404)
        conn.execute(
            "UPDATE deliveries SET status=? WHERE id=?",
            (status, int(m.group(1))),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM deliveries WHERE id=?", (int(m.group(1)),)).fetchone()
        conn.close()
        return self.send_json(row_to_dict(row))

    def do_DELETE(self):
        if not self.trusted_origin():
            return self.send_json({"error": "Origine non autorisee"}, 403)
        if not self.require_editor():
            return
        order_match = re.fullmatch(r"/api/deliveries/(\d+)/order-files/(\d+)", urlparse(self.path).path)
        if order_match:
            if self.session().get("username") != "admin":
                return self.send_json({"error": "Suppression réservée au compte admin"}, 403)
            delivery_id, file_id = map(int, order_match.groups())
            with closing(db()) as conn, conn:
                file = conn.execute("SELECT id FROM order_files WHERE delivery_id=? AND id=?",
                    (delivery_id, file_id)).fetchone()
                if not file:
                    return self.send_json({"error": "Bon de commande introuvable"}, 404)
                history = conn.execute("SELECT 1 FROM order_item_history WHERE file_id=? LIMIT 1", (file_id,)).fetchone()
                if history:
                    return self.send_json({"error": "Ce bon contient un historique de chargement et ne peut pas être supprimé."}, 409)
                conn.execute("DELETE FROM order_files WHERE delivery_id=? AND id=?", (delivery_id, file_id))
            return self.send_json({"ok": True})
        m = re.fullmatch(r"/api/deliveries/(\d+)", urlparse(self.path).path)
        if not m:
            return self.send_json({"error": "Not found"}, 404)
        conn = db()
        cur = conn.execute("DELETE FROM deliveries WHERE id=?", (int(m.group(1)),))
        conn.commit()
        conn.close()
        if cur.rowcount == 0:
            return self.send_json({"error": "Introuvable"}, 404)
        return self.send_json({"ok": True})


SHEET_META = {
    "Planning": ("Planning", "EN ATTENTE"),
    "TECLIM": ("TECLIM", "EN ATTENTE"),
    "MV EN COURS DE LIVRAISON": ("ops", "EN COURS DE LIVRAISON"),
    "MV LIVRÉ": ("ops", "LIVRÉ"),
    "MV LIVRÉ ": ("ops", "LIVRÉ"),
}
CLE_FILL = {
    "FFD9EAD3": "LIVRÉ",
    "D9EAD3": "LIVRÉ",
    "FFFCE5CD": "EN COURS DE LIVRAISON",
    "FCE5CD": "EN COURS DE LIVRAISON",
}


def cell_cle_status(cell) -> str:
    """Map Excel fill to the CLÉ sheet: green = LIVRÉ, orange = EN COURS."""
    fill = getattr(cell, "fill", None)
    if not fill or fill.patternType in (None, "none"):
        return ""
    fg = fill.fgColor
    rgb = ""
    if fg is not None and getattr(fg, "rgb", None):
        rgb = str(fg.rgb).upper()
    if rgb in ("00000000", "FFFFFFFF", "0"):
        return ""
    if rgb in CLE_FILL:
        return CLE_FILL[rgb]
    if rgb.endswith("D9EAD3"):
        return "LIVRÉ"
    if rgb.endswith("FCE5CD"):
        return "EN COURS DE LIVRAISON"
    return ""


def excel_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute:
            return value.strftime("%d/%m/%Y %H:%M")
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)) and 20000 < float(value) < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def excel_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20000 < float(value) < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date().isoformat()
    parsed = parse_date(str(value).strip())
    return parsed or ""


def import_workbook(path: str, replace: bool = True) -> int:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    records = []
    seen: set[tuple] = set()
    for sheet_name in wb.sheetnames:
        meta = SHEET_META.get(sheet_name) or SHEET_META.get(sheet_name.strip())
        if not meta:
            continue
        category, default_status = meta
        ws = wb[sheet_name]
        header_row = None
        col = {}
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=6, values_only=True), 1):
            labels = [excel_text(c).upper() for c in row]
            if labels and labels[0] in ("", " "):
                labels[0] = "NO"
            if "VESSEL" in labels:
                header_row = i
                col = {lab: idx for idx, lab in enumerate(labels) if lab}
                break
        if not header_row:
            continue

        def take(vals, *names):
            for name in names:
                idx = col.get(name)
                if idx is not None and idx < len(vals):
                    return vals[idx]
            return None

        for row in ws.iter_rows(min_row=header_row + 1):
            vals = [c.value for c in row]
            vessel = excel_text(take(vals, "VESSEL"))
            if not vessel or vessel.upper() in {"VESSEL", "NO"}:
                continue
            status_text = excel_text(take(vals, "STATUS"))
            # MV LIVRÉ sheet + green CLÉ = LIVRÉ, even if the STATUS cell still says EN ATTENTE
            if default_status == "LIVRÉ":
                status = "LIVRÉ"
            elif status_text:
                status = status_text
            else:
                idx = col.get("STATUS")
                color_status = cell_cle_status(row[idx]) if idx is not None and idx < len(row) else ""
                # TECLIM STATUS column is pre-filled green — not a per-row CLÉ
                if category == "TECLIM" and color_status == "LIVRÉ":
                    color_status = ""
                status = color_status or default_status
            rec = {
                "no": take(vals, "NO", " "),
                "vessel": vessel,
                "imo": excel_text(take(vals, "IMO")),
                "port": excel_text(take(vals, "PORT")),
                "agent": excel_text(take(vals, "AGENT")),
                "etb": excel_text(take(vals, "ETB")),
                "etd": excel_text(take(vals, "ETD")),
                "quai": excel_text(take(vals, "QUAI")),
                "commande": excel_text(take(vals, "COMMANDE")),
                "status": status,
                "date_of_delivery": excel_date(take(vals, "DATE OF DELIVERY")),
                "livreur": excel_text(take(vals, "LIVREUR")),
                "invoicing_date": excel_date(take(vals, "INVOICING DATE")),
                "category": category,
                "year": None,
            }
            try:
                rec = payload_to_record(rec)
            except ValueError:
                continue
            key = (
                rec["vessel"].upper(),
                rec["port"],
                rec["date_of_delivery"],
                rec["status"],
                rec["category"],
            )
            if key in seen:
                continue
            seen.add(key)
            records.append(rec)

    init_db()
    conn = db()
    if replace:
        conn.execute("DELETE FROM deliveries")
    for rec in records:
        conn.execute(
            """
            INSERT INTO deliveries (
                no, vessel, imo, port, agent, etb, etd, quai, commande, status,
                date_of_delivery, livreur, invoicing_date, category, year
            ) VALUES (:no,:vessel,:imo,:port,:agent,:etb,:etd,:quai,:commande,:status,
                      :date_of_delivery,:livreur,:invoicing_date,:category,:year)
            """,
            rec,
        )
    conn.commit()
    n = conn.execute("SELECT COUNT(*) FROM deliveries").fetchone()[0]
    conn.close()
    return n


def backup_dirs() -> list[str]:
    dirs = []
    local = (BACKUP_LOCAL or "").strip()
    if local:
        dirs.append(local)
    share = (BACKUP_SHARE or "").strip()
    if share and (os.path.isdir(os.path.dirname(share)) or os.path.isdir(r"E:\Partage")):
        dirs.append(share)
    return dirs


def prune_backups(folder: str, keep: int = 14) -> None:
    files = sorted(glob.glob(os.path.join(folder, "marine_ops_*.sqlite")))
    for old in files[:-keep]:
        try:
            os.remove(old)
        except OSError:
            pass


def copy_db(dest_path: str) -> None:
    os.makedirs(os.path.dirname(dest_path), exist_ok=True)
    src = sqlite3.connect(DB_PATH)
    dst = sqlite3.connect(dest_path)
    with dst:
        src.backup(dst)
    dst.close()
    src.close()


def backup_if_needed(force: bool = False) -> str | None:
    if not os.path.isfile(DB_PATH):
        return None
    name = f"marine_ops_{today().isoformat()}.sqlite"
    written = None
    for folder in backup_dirs():
        try:
            os.makedirs(folder, exist_ok=True)
            dest = os.path.join(folder, name)
            if not force and os.path.isfile(dest):
                if os.path.getmtime(dest) >= os.path.getmtime(DB_PATH) - 2:
                    prune_backups(folder)
                    continue
            copy_db(dest)
            prune_backups(folder)
            written = dest
            print(f"sauvegarde → {dest}")
        except OSError as exc:
            print(f"sauvegarde ignoree ({folder}): {exc}")
    return written


def backup_loop() -> None:
    while True:
        time.sleep(3600)
        try:
            backup_if_needed()
        except Exception as exc:
            print(f"sauvegarde: {exc}")


def listen_ports() -> list[int]:
    flag = os.environ.get("PORT_STRICT", "").strip().lower()
    if flag in ("1", "true", "yes"):
        return [PORT]
    return list(dict.fromkeys([PORT, 80, 3847]))


def main() -> None:
    auth_account()
    init_db()
    try:
        backup_if_needed(force=True)
    except Exception as exc:
        print(f"sauvegarde: {exc}")
    threading.Thread(target=backup_loop, daemon=True).start()
    tried = []
    httpd = None
    bound = PORT
    for p in listen_ports():
        tried.append(p)
        try:
            httpd = ThreadingHTTPServer((HOST, p), Handler)
            bound = p
            break
        except OSError:
            httpd = None
    if httpd is None:
        raise SystemExit(f"Impossible d'ouvrir le port ({tried})")
    extra = "" if bound in (80, 443) else f":{bound}"
    print(f"MANYANI SHIPPING → http://{SITE}{extra}")
    print(f"MANYANI SHIPPING → http://{LAN_IP}{extra}")
    print(f"base → {DB_PATH}")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("\nstop")
        httpd.server_close()


if __name__ == "__main__":
    main()
