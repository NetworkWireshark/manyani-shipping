"""Read purchase-order line items from XLS and XLSX workbooks."""
import io
import math
import re
import unicodedata


# First matching column wins. Bare DISCOUNT is a currency amount on AVS sheets
# and must not replace DISC% (the percent).
_ALIASES = {
    "NO": "no", "N": "no", "NUMBER": "no", "NUMERO": "no",
    "CODE": "code",
    "DESCRIPTION": "description", "DESIGNATION": "description", "ARTICLE": "description",
    "ARTICLE / DESCRIPTION": "description",
    "MAKER/BRAND": "brand", "MARQUE": "brand", "BRAND": "brand",
    "REMARK": "remark", "REMARQUE": "remark",
    "UNIT": "unit", "UNITE": "unit", "UNITY": "unit",
    "QTY": "quantity", "QUANTITY": "quantity", "QUANTITE": "quantity", "QTE": "quantity",
    "DISC%": "discount", "REMISE": "discount", "REMISE%": "discount", "DISCOUNT%": "discount",
    "PRICE": "price", "UNIT PRICE": "price", "PRIX UNITAIRE": "price", "PRIX": "price",
    "AMOUNT": "amount", "MONTANT": "amount",
    "NET PRICE": "net_price", "PRIX NET": "net_price",
    "NETTOTAL": "net_total", "NET TOTAL": "net_total", "TOTAL NET": "net_total",
}
_CURRENCY = {"USD": "USD", "EUR": "EUR", "EURO": "EUR", "MAD": "MAD"}
_META = {
    "VESSEL NAME": "vessel", "DELIVERY PORT": "port", "ETA/DELIVERY DATE": "delivery_date",
    "CURRENCY": "currency", "TOTAL AMOUNT": "total", "GRAND TOTAL": "grand_total", "REF NO": "reference",
}
_MONEY = ("price", "discount", "amount", "net_price", "net_total")


def parse_order(content):
    if content.startswith(b"%PDF-"):
        from order_pdf import parse_pdf_order
        return parse_pdf_order(content)
    if content.startswith(bytes.fromhex("d0cf11e0a1b11ae1")):
        import xlrd
        workbook = xlrd.open_workbook(file_contents=content)
        sheets = [(sheet.name, [sheet.row_values(i) for i in range(sheet.nrows)]) for sheet in workbook.sheets()]
        workbook.release_resources()
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheets = [(sheet.title, list(sheet.iter_rows(values_only=True))) for sheet in workbook]
        finally:
            workbook.close()
    return parse_sheets(sheets)


def _header_text(value):
    text = " ".join(str(value or "").replace("\xa0", " ").split()).upper()
    text = "".join(c for c in unicodedata.normalize("NFD", text) if not unicodedata.combining(c))
    text = text.replace("°", "").replace("º", "").replace("(", " ").replace(")", " ")
    return " ".join(text.split())


def _column_name(value):
    text = _header_text(value)
    if not text:
        return None, None
    match = re.fullmatch(r"TOTAL IN (USD|EUR|EURO|MAD)", text)
    if match:
        return "amount", _CURRENCY[match.group(1)]
    compact = text.replace(" %", "%")
    key = _ALIASES.get(compact) or _ALIASES.get(text)
    if key:
        return key, None
    if text.startswith("PRIX UNITAIRE") or text.startswith("UNIT PRICE"):
        return "price", None
    if text.startswith("TOTAL NET") or text.startswith("NET TOTAL"):
        return "net_total", None
    if text.startswith("PRIX NET") or text.startswith("NET PRICE"):
        return "net_price", None
    return None, None


def _map_columns(row):
    columns = {}
    currency = None
    for column, value in enumerate(row):
        key, hint = _column_name(value)
        if key and key not in columns:
            columns[key] = column
        if hint and currency is None:
            currency = hint
    return columns, currency


def _money(number):
    if number is None:
        return None
    return round(float(number), 6)


def _number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    text = str(value or "").strip().replace("\u00a0", "").replace("\u202f", "").replace(" ", "").replace("%", "")
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    if not re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        return None
    result = float(text)
    return result if math.isfinite(result) else None


def _complete_money(item):
    # Keep numbers already present on the row. Fill only the gaps.
    qty = _number(item.get("quantity"))
    if qty is None:
        return False
    item["quantity"] = qty
    price = _number(item.get("price"))
    amount = _number(item.get("amount"))
    discount = _number(item.get("discount"))
    net_price = _number(item.get("net_price"))
    had_net = "net_total" in item
    net_total = _number(item.get("net_total")) if had_net else None
    price = _money(price)
    amount = _money(amount)
    discount = _money(discount)
    net_price = _money(net_price)
    net_total = _money(net_total)
    if price is None and amount is not None and qty:
        price = round(amount / qty, 6)
    if amount is None and price is not None:
        amount = round(float(price) * float(qty), 6)
    if discount is None:
        discount = 0
    if price is not None:
        item["price"] = price
    if amount is not None:
        item["amount"] = amount
    item["discount"] = discount
    if net_price is None and price is not None:
        net_price = round(float(price) * (1 - float(discount) / 100), 6)
    if net_price is not None:
        item["net_price"] = net_price
    if net_total is not None:
        item["net_total"] = net_total
    elif had_net:
        item.pop("net_total", None)
    if "net_total" not in item and discount not in (0, 0.0) and price is not None:
        item["net_total"] = round(float(price) * (1 - float(discount) / 100) * float(qty), 6)
    return True


def _sum_items(items, key):
    total = 0.0
    for item in items:
        number = _number(item.get(key))
        if number is None:
            return None
        total += float(number)
    return total


def parse_sheets(sheets):
    metadata = {}
    items = []
    groups = []
    for sheet_name, rows in sheets:
        header = None
        columns = None
        currency_hint = None
        for index, row in enumerate(rows):
            mapped, hint = _map_columns(row)
            if "description" in mapped and "quantity" in mapped:
                header = index
                columns = mapped
                currency_hint = hint
                break
        if header is None:
            continue
        for row in rows[:header]:
            for column, value in enumerate(row):
                key = _META.get(_header_text(value))
                if not key:
                    continue
                following = next((cell for cell in row[column + 1:] if cell is not None and str(cell).strip()), "")
                metadata[key] = following
        if currency_hint and not str(metadata.get("currency") or "").strip():
            metadata["currency"] = currency_hint
        group = ""
        for row in rows[header + 1:]:
            item = {}
            for key, column in columns.items():
                raw = row[column] if column < len(row) and row[column] is not None else ""
                if key in _MONEY or key == "quantity":
                    number = _number(raw)
                    item[key] = number if number is not None else ""
                else:
                    item[key] = raw.strip() if isinstance(raw, str) else raw
            item["description"] = str(item.get("description", "")).strip()
            if not item["description"]:
                continue
            if str(item.get("no", "")).strip().upper() == "G":
                group = item["description"]
                groups.append(group)
                continue
            if not _complete_money(item):
                continue
            item["group"] = group
            item["sheet"] = sheet_name
            items.append(item)
    if _number(metadata.get("total")) is None:
        amount_sum = _sum_items(items, "amount")
        if amount_sum is not None:
            metadata["total"] = round(amount_sum, 6)
    if _number(metadata.get("grand_total")) is None:
        net_sum = _sum_items(items, "net_total")
        if net_sum is None:
            net_sum = _sum_items(items, "amount")
        if net_sum is not None:
            metadata["grand_total"] = round(net_sum, 6)
    if not items:
        raise ValueError("Aucun article trouve. Colonnes Description et Quantite requises.")
    return {"metadata": metadata, "items": items, "groups": groups}
